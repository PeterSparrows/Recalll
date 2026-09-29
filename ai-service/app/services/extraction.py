"""
Extraction service: pulls raw text (plus page/slide counts) out of
PDF, DOCX, PPTX, and TXT files.

All libraries here (PyMuPDF, pdfplumber, python-docx, python-pptx) run
100% locally with no network calls or model downloads.
"""
from dataclasses import dataclass, field
from typing import List
import fitz  # PyMuPDF
import pdfplumber
from docx import Document as DocxDocument
from pptx import Presentation


@dataclass
class ExtractedPage:
    page_number: int
    text: str


@dataclass
class ExtractionResult:
    pages: List[ExtractedPage] = field(default_factory=list)
    page_count: int = 0

    @property
    def full_text(self) -> str:
        return "\n\n".join(p.text for p in self.pages if p.text.strip())


class UnsupportedFileTypeError(Exception):
    pass


class ExtractionFailedError(Exception):
    pass


def extract_pdf(file_path: str) -> ExtractionResult:
    """
    Extracts text page-by-page from a PDF. Uses PyMuPDF as the primary
    engine (fast, robust); falls back to pdfplumber per-page if a page
    comes back empty (some PDFs extract better with one library than
    the other, especially those with unusual encodings).
    """
    pages: List[ExtractedPage] = []
    try:
        doc = fitz.open(file_path)
        for i, page in enumerate(doc):
            text = page.get_text("text") or ""
            if not text.strip():
                text = _pdfplumber_fallback_page(file_path, i)
            pages.append(ExtractedPage(page_number=i + 1, text=text))
        doc.close()
    except Exception as e:
        raise ExtractionFailedError(f"Failed to extract PDF: {e}") from e

    if not pages:
        raise ExtractionFailedError("PDF appears to have no extractable pages.")

    return ExtractionResult(pages=pages, page_count=len(pages))


def _pdfplumber_fallback_page(file_path: str, index: int) -> str:
    try:
        with pdfplumber.open(file_path) as pdf:
            if index < len(pdf.pages):
                return pdf.pages[index].extract_text() or ""
    except Exception:
        pass
    return ""


def extract_docx(file_path: str) -> ExtractionResult:
    """
    Extracts text from a .docx file, paragraph by paragraph, including
    text inside tables. DOCX has no native "page" concept (pagination
    is a rendering detail), so we treat the whole document as a single
    logical page for chunking purposes downstream.
    """
    try:
        doc = DocxDocument(file_path)
        parts: List[str] = []

        for para in doc.paragraphs:
            if para.text.strip():
                parts.append(para.text)

        for table in doc.tables:
            for row in table.rows:
                row_text = " | ".join(cell.text.strip() for cell in row.cells if cell.text.strip())
                if row_text:
                    parts.append(row_text)

        text = "\n".join(parts)
    except Exception as e:
        raise ExtractionFailedError(f"Failed to extract DOCX: {e}") from e

    if not text.strip():
        raise ExtractionFailedError("DOCX appears to contain no extractable text.")

    return ExtractionResult(pages=[ExtractedPage(page_number=1, text=text)], page_count=1)


def extract_pptx(file_path: str) -> ExtractionResult:
    """Extracts text slide-by-slide, including speaker notes."""
    pages: List[ExtractedPage] = []
    try:
        prs = Presentation(file_path)
        for i, slide in enumerate(prs.slides):
            parts: List[str] = []
            for shape in slide.shapes:
                if shape.has_text_frame:
                    for para in shape.text_frame.paragraphs:
                        line = "".join(run.text for run in para.runs)
                        if line.strip():
                            parts.append(line)
                if shape.has_table:
                    for row in shape.table.rows:
                        row_text = " | ".join(cell.text.strip() for cell in row.cells if cell.text.strip())
                        if row_text:
                            parts.append(row_text)
            if slide.has_notes_slide and slide.notes_slide.notes_text_frame:
                notes = slide.notes_slide.notes_text_frame.text
                if notes.strip():
                    parts.append(f"[Speaker notes: {notes.strip()}]")
            pages.append(ExtractedPage(page_number=i + 1, text="\n".join(parts)))
    except Exception as e:
        raise ExtractionFailedError(f"Failed to extract PPTX: {e}") from e

    if not pages:
        raise ExtractionFailedError("PPTX appears to have no slides.")

    return ExtractionResult(pages=pages, page_count=len(pages))


def extract_txt(file_path: str) -> ExtractionResult:
    """Reads a plain text file, trying a couple of common encodings."""
    text = None
    for encoding in ("utf-8", "utf-8-sig", "latin-1"):
        try:
            with open(file_path, "r", encoding=encoding) as f:
                text = f.read()
            break
        except UnicodeDecodeError:
            continue
        except Exception as e:
            raise ExtractionFailedError(f"Failed to read TXT: {e}") from e

    if text is None:
        raise ExtractionFailedError("Could not decode text file with any supported encoding.")
    if not text.strip():
        raise ExtractionFailedError("TXT file is empty.")

    return ExtractionResult(pages=[ExtractedPage(page_number=1, text=text)], page_count=1)


EXTRACTORS = {
    "pdf": extract_pdf,
    "docx": extract_docx,
    "pptx": extract_pptx,
    "txt": extract_txt,
}


def extract(file_path: str, file_type: str) -> ExtractionResult:
    file_type = file_type.lower().lstrip(".")
    extractor = EXTRACTORS.get(file_type)
    if extractor is None:
        raise UnsupportedFileTypeError(f"Unsupported file type: {file_type}")
    return extractor(file_path)
